from datetime import datetime, timedelta, timezone
from io import BytesIO

import pandas as pd
import streamlit as st
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def create_excel(data, filters, report_name, username):
    buffer = BytesIO()

    vietnam_time = datetime.now(
        timezone(timedelta(hours=7))
    ).strftime("%d/%m/%Y %H:%M:%S")

    metadata = [
        ("Tên báo cáo", report_name),
        ("Nguồn dữ liệu", "Olist Brazil"),
        ("Đơn vị tiền", "BRL"),
        ("Người xuất", username),
        ("Thời điểm tạo file (UTC+7)", vietnam_time),
        ("Số dòng dữ liệu", len(data)),
        ("Phạm vi", "Toàn bộ kết quả phù hợp bộ lọc"),
    ]

    metadata.extend(
        (str(name), str(value))
        for name, value in filters.items()
    )

    info = pd.DataFrame(
        metadata,
        columns=["Thông tin", "Giá trị"],
    )

    with pd.ExcelWriter(
        buffer,
        engine="openpyxl",
        datetime_format="DD/MM/YYYY HH:MM:SS",
        date_format="DD/MM/YYYY",
    ) as writer:
        data.to_excel(
            writer,
            sheet_name="Du_lieu",
            index=False,
        )

        info.to_excel(
            writer,
            sheet_name="Thong_tin",
            index=False,
        )

        for sheet in writer.book.worksheets:
            sheet.freeze_panes = "A2"

            if sheet.title == "Du_lieu":
                sheet.auto_filter.ref = sheet.dimensions

            # Định dạng dòng tiêu đề
            for cell in sheet[1]:
                cell.fill = PatternFill(
                    fill_type="solid",
                    fgColor="1D4ED8",
                )
                cell.font = Font(
                    bold=True,
                    color="FFFFFF",
                )
                cell.alignment = Alignment(
                    vertical="center",
                )

            sheet.row_dimensions[1].height = 26

            # Tính chiều rộng dựa trên tối đa 200 dòng dữ liệu
            for column_index in range(1, sheet.max_column + 1):
                values = [
                    sheet.cell(row, column_index).value
                    for row in range(
                        1,
                        min(sheet.max_row, 201) + 1,
                    )
                ]

                max_length = max(
                    len(str(value)) if value is not None else 0
                    for value in values
                )

                letter = get_column_letter(column_index)

                sheet.column_dimensions[letter].width = min(
                    max(max_length + 2, 14),
                    55,
                )

            # Giữ chuỗi bắt đầu bằng "=" dưới dạng văn bản
            for row in sheet.iter_rows(min_row=2):
                for cell in row:
                    if cell.data_type == "f":
                        cell.data_type = "s"

        info_sheet = writer.book["Thong_tin"]
        info_sheet.column_dimensions["A"].width = 30
        info_sheet.column_dimensions["B"].width = 75

        for row in info_sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(
                    vertical="top",
                    wrap_text=True,
                )

    return buffer.getvalue()


def render_excel_export(
    data,
    filters,
    report_name,
    username,
    file_name,
    key,
):
    if data.empty:
        return

    # Chỉ tạo file khi người dùng yêu cầu.
    if st.button(
        "Chuẩn bị file Excel",
        key=f"{key}_prepare",
    ):
        try:
            with st.spinner("Đang tạo file Excel..."):
                excel_bytes = create_excel(
                    data=data,
                    filters=filters,
                    report_name=report_name,
                    username=username,
                )
        except Exception as error:
            st.error(f"Không tạo được Excel: {error}")
            return

        st.download_button(
            label="Tải file Excel",
            data=excel_bytes,
            file_name=file_name,
            mime=(
                "application/vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            ),
            key=f"{key}_download",
        )

        st.caption(
            f"File có {len(data):,} dòng dữ liệu. "
            "Sau khi đổi bộ lọc, hãy chuẩn bị file mới."
        )
